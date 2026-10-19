"""Real file exporters (optional deps: python-pptx, python-docx). If a library is missing the
exporter returns None and the caller records that the file was NOT produced."""
from __future__ import annotations

import io
import json
import re
from typing import Optional


def parse_json_loose(text: str):
    """Parse JSON from model output, tolerating code fences and surrounding prose."""
    t = text.strip()
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t, flags=re.S).strip()
    try:
        return json.loads(t)
    except ValueError:
        start, end = t.find("{"), t.rfind("}")
        if start != -1 and end > start:
            return json.loads(t[start:end + 1])
        raise


def build_pptx(outline: dict) -> Optional[bytes]:
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError:
        return None
    slides = outline.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("outline has no slides")
    prs = Presentation()
    for i, s in enumerate(slides[:40]):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(s.get("title", ""))[:200]
        bullets = [str(b)[:300] for b in (s.get("bullets") or [])][:10]
        placeholder = slide.placeholders[1]
        if i == 0:
            placeholder.text = " | ".join(bullets)[:300]
        else:
            tf = placeholder.text_frame
            tf.text = bullets[0] if bullets else ""
            for b in bullets[1:]:
                tf.add_paragraph().text = b
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(22)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])[:2000]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_docx(markdown: str) -> Optional[bytes]:
    try:
        from docx import Document
    except ImportError:
        return None
    doc = Document()
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("> "):
            doc.add_paragraph(line[2:]).italic = True
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            doc.add_paragraph(re.sub(r"\*\*(.*?)\*\*", r"\1", line))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
